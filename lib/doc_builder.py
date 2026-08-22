"""
doc_builder.py — Step 4: Generate complete Word document.
=========================================================================
Uses python-docx to assemble a full confined space construction plan.
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
import json


def _set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_heading(doc, text, level=1, color="1F4E79"):
    """Add a styled heading."""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    return h


def _add_table(doc, headers, rows, col_widths=None):
    """Add a formatted table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        _set_cell_shading(cell, "1F4E79")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = str(value) if value is not None else ""
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
            if r % 2 == 1:
                _set_cell_shading(cell, "F2F2F2")

    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Inches(width)
    return table


def generate_plan_docx(plan_data: dict, output_path: str) -> str:
    """Generate complete confined space construction plan as Word document.

    Args:
        plan_data: dict with all plan sections (project_info, space_description,
                   work_content, hazard_assessment, control_measures,
                   permit, emergency_plan, compliance_report)
        output_path: Path to save .docx file

    Returns:
        Path to saved file
    """
    doc = Document()

    # --- Page setup ---
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    # --- Cover Page ---
    project = plan_data.get("project_info", {})
    space = plan_data.get("space_description", {})
    work = plan_data.get("work_content", {})

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("\n\n\n")
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CONFINED SPACE\nCONSTRUCTION PLAN")
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("1F4E79")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"\n{project.get('name', '[Project Name]')}")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string("333333")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"\n{project.get('contractor', '[Contractor Name]')}")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string("666666")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    permit_id = f"LOTO-CS-{datetime.now().strftime('%Y%m%d')}-001"
    run = p.add_run(f"\nPermit No: {permit_id}")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("999999")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"\nDate: {datetime.now().strftime('%Y-%m-%d')}\nVersion: 1.0\n\n")
    run.font.size = Pt(11)

    doc.add_page_break()

    # --- Table of Contents ---
    _add_heading(doc, "Table of Contents", level=1)
    toc_items = [
        "1. Regulatory Basis",
        "2. Project Overview",
        "3. Risk Assessment Report",
        "4. Safety Measures",
        "5. Work Permit",
        "6. Emergency Plan",
        "7. Attachments",
        "8. Compliance Review Report",
    ]
    for item in toc_items:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.space_after = Pt(4)

    doc.add_page_break()

    # --- Chapter 1: Regulatory Basis ---
    _add_heading(doc, "1. Regulatory Basis", level=1)
    _add_heading(doc, "1.1 Applicable Regulations", level=2)
    _add_table(doc,
        ["Regulation", "Purpose"],
        [
            ["Safety Regulations §161-172", "Statutory definition, risk assessment, permit, 17 safety measures"],
            ["Safety Management Manual A.2.5/C.1.8", "Project-level standards, competent person, inspection checklist"],
            ["ISO 31000", "Risk management framework, risk matrix, hazard identification"],
        ],
        col_widths=[2.5, 4.0])

    _add_heading(doc, "1.2 Statutory Definition of Confined Space", level=2)
    doc.add_paragraph(
        "A confined space is any enclosed space that may present one or more of the following 6 dangers:"
    )
    dangers = plan_data.get("hazard_assessment", {}).get("legal_definition", [])
    for i, danger in enumerate(dangers, 1):
        doc.add_paragraph(f"{i}. {danger}", style="List Number")

    doc.add_page_break()

    # --- Chapter 2: Project Overview ---
    _add_heading(doc, "2. Project Overview", level=1)
    _add_heading(doc, "2.1 Project Information", level=2)
    _add_table(doc,
        ["Item", "Details"],
        [
            ["Project Name", project.get("name", "")],
            ["Location", project.get("location", "")],
            ["Contractor", project.get("contractor", "")],
            ["Permit Number", permit_id],
        ],
        col_widths=[2.0, 4.5])

    _add_heading(doc, "2.2 Confined Space Description", level=2)
    _add_table(doc,
        ["Parameter", "Description"],
        [
            ["Space Type", space.get("type", "")],
            ["Dimensions", space.get("dimensions", "")],
            ["Location", space.get("location", "")],
            ["Access Points", space.get("access", "")],
            ["Internal Structure", space.get("structure", "")],
            ["Surroundings", space.get("surroundings", "")],
        ],
        col_widths=[2.0, 4.5])

    _add_heading(doc, "2.3 Work Content", level=2)
    _add_table(doc,
        ["Parameter", "Description"],
        [
            ["Work Nature", work.get("nature", "")],
            ["Number of Workers", str(work.get("workers", ""))],
            ["Duration", work.get("duration", "")],
            ["Materials/Equipment", work.get("materials", "")],
        ],
        col_widths=[2.0, 4.5])

    doc.add_page_break()

    # --- Chapter 3: Risk Assessment ---
    _add_heading(doc, "3. Risk Assessment Report", level=1)

    hazards = plan_data.get("hazard_assessment", {}).get("hazards", [])
    _add_heading(doc, "3.1 Identified Hazards", level=2)
    _add_table(doc,
        ["Hazard", "Category", "Source", "Severity", "Likelihood", "Risk Score", "Risk Level"],
        [[h.get("name", h.get("key", "")), h.get("category", ""), h.get("source", ""),
          f"{h.get('severity', 3)} ({h.get('severity_label', '')})",
          f"{h.get('likelihood', 3)} ({h.get('likelihood_label', '')})",
          str(h.get("risk_score", "")),
          h.get("risk_level", "")]
         for h in hazards],
        col_widths=[1.2, 0.8, 1.2, 0.8, 0.8, 0.5, 0.7])

    _add_heading(doc, "3.2 Risk Assessment Method", level=2)
    doc.add_paragraph("Risk Score = Severity × Likelihood (ISO 31000 framework)")
    _add_table(doc,
        ["Risk Level", "Score Range", "Action Required"],
        [
            ["Extreme", "15-25", "Do not proceed. Must eliminate risk before re-assessment."],
            ["High", "8-14", "Requires advanced control measures + project manager approval."],
            ["Medium", "4-7", "Requires control measures + attendant monitoring."],
            ["Low", "1-3", "Standard PPE + attendant monitoring."],
        ],
        col_widths=[1.2, 1.0, 4.3])

    if "ventilation_requirement" in plan_data.get("hazard_assessment", {}):
        vent = plan_data["hazard_assessment"]["ventilation_requirement"]
        _add_heading(doc, "3.3 Ventilation Calculation", level=2)
        _add_table(doc,
            ["Parameter", "Value"],
            [
                ["Space Volume", f"{vent['space_volume_m3']} m³"],
                ["Required Airflow", f"{vent['required_airflow_m3h']} m³/h"],
                ["Air Changes per Hour", f"{vent['air_changes_per_hour']}"],
                ["Minimum Ventilation Time", f"{vent['min_ventilation_time_min']} min"],
                ["Recommended Equipment", vent["recommended_equipment"]],
            ],
            col_widths=[2.5, 4.0])

    doc.add_page_break()

    # --- Chapter 4: Safety Measures ---
    _add_heading(doc, "4. Safety Measures", level=1)

    controls = plan_data.get("control_measures", [])
    _add_heading(doc, "4.1 Engineering Controls", level=2)
    for ctrl in controls:
        if ctrl.get("engineering"):
            p = doc.add_paragraph()
            run = p.add_run(f"Hazard: {ctrl.get('hazard_key', '').replace('_', ' ').title()}")
            run.font.bold = True
            for measure in ctrl["engineering"]:
                doc.add_paragraph(measure, style="List Bullet")

    _add_heading(doc, "4.2 Administrative Controls", level=2)
    for ctrl in controls:
        if ctrl.get("administrative"):
            p = doc.add_paragraph()
            run = p.add_run(f"Hazard: {ctrl.get('hazard_key', '').replace('_', ' ').title()}")
            run.font.bold = True
            for measure in ctrl["administrative"]:
                doc.add_paragraph(measure, style="List Bullet")

    _add_heading(doc, "4.3 Personal Protective Equipment (PPE)", level=2)
    all_ppe = set()
    for ctrl in controls:
        for ppe in ctrl.get("ppe", []):
            all_ppe.add(ppe)
    for ppe in sorted(all_ppe):
        doc.add_paragraph(ppe, style="List Bullet")

    _add_heading(doc, "4.4 Personnel Requirements", level=2)
    from .control_generator import get_personnel_requirements
    personnel = get_personnel_requirements()
    _add_table(doc,
        ["Role", "Requirement", "Quantity"],
        [[p["role"], p["requirement"], p["quantity"]] for p in personnel],
        col_widths=[1.5, 3.5, 1.0])

    doc.add_page_break()

    # --- Chapter 5: Work Permit ---
    _add_heading(doc, "5. Work Permit", level=1)

    _add_heading(doc, "5.1 Permit Information", level=2)
    _add_table(doc,
        ["Item", "Details"],
        [
            ["Permit Number", permit_id],
            ["Work Content", work.get("nature", "")],
            ["Valid From", datetime.now().strftime("%Y-%m-%d %H:%M")],
            ["Valid Until", "(Max one shift)"],
            ["Issuer (Competent Person)", "_____________________"],
            ["Receiver (Authorized Entrant)", "_____________________"],
        ],
        col_widths=[2.5, 4.0])

    _add_heading(doc, "5.2 Permit Conditions Checklist (12 items)", level=2)
    from .control_generator import get_permit_conditions
    conditions = get_permit_conditions()
    for i, cond in enumerate(conditions, 1):
        p = doc.add_paragraph()
        run = p.add_run(f"☐  {i}. {cond}")
        run.font.size = Pt(10)

    _add_heading(doc, "5.3 Statutory Safety Measures (17 items)", level=2)
    from .control_generator import get_safety_measures
    measures = get_safety_measures()
    for i, measure in enumerate(measures, 1):
        p = doc.add_paragraph()
        run = p.add_run(f"☐  {i}. {measure}")
        run.font.size = Pt(10)

    doc.add_page_break()

    # --- Chapter 6: Emergency Plan ---
    _add_heading(doc, "6. Emergency Plan", level=1)

    _add_heading(doc, "6.1 Rescue Equipment", level=2)
    from .control_generator import get_rescue_equipment
    rescue_eq = get_rescue_equipment()
    _add_table(doc,
        ["Equipment", "Quantity", "Purpose"],
        [[e["item"], e["quantity"], e["purpose"]] for e in rescue_eq],
        col_widths=[2.0, 1.5, 3.0])

    _add_heading(doc, "6.2 Rescue Procedure", level=2)
    steps = [
        "Attendant detects abnormality → DO NOT enter the space",
        "Immediately notify site supervisor + call emergency number",
        "Use tripod + winch to attempt external extraction",
        "If external rescue fails → rescue personnel enter with SCBA",
        "After extraction, immediately administer first aid / CPR",
        "Notify emergency services (999 or local equivalent)",
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f"Step {i}: {step}", style="List Number")

    _add_heading(doc, "6.3 Emergency Contacts", level=2)
    _add_table(doc,
        ["Service", "Number"],
        [
            ["Fire/Ambulance", "999 (or local emergency number)"],
            ["Police", "999 (or local emergency number)"],
            ["Labor Bureau", "[Local labor bureau number]"],
            ["Project Emergency Contact", "[Site emergency number]"],
        ],
        col_widths=[3.0, 3.5])

    doc.add_page_break()

    # --- Chapter 7: Attachments ---
    _add_heading(doc, "7. Attachments", level=1)

    _add_heading(doc, "7.1 Gas Monitoring Record", level=2)
    _add_table(doc,
        ["Time", "Position", "O₂(%)", "CO(ppm)", "H₂S(ppm)", "LEL(%)", "Inspector", "Remarks"],
        [["", "Top", "", "", "", "", "", ""],
         ["", "Middle", "", "", "", "", "", ""],
         ["", "Bottom", "", "", "", "", "", ""]],
        col_widths=[0.8, 0.7, 0.6, 0.7, 0.7, 0.6, 0.8, 0.8])

    _add_heading(doc, "7.2 Entry/Exit Log", level=2)
    _add_table(doc,
        ["Name", "ID No.", "Entry Time", "Exit Time", "Attendant Confirm", "Remarks"],
        [["", "", "", "", "", ""],
         ["", "", "", "", "", ""],
         ["", "", "", "", "", ""]],
        col_widths=[1.2, 1.0, 1.0, 1.0, 1.2, 1.1])

    _add_heading(doc, "7.3 LOTO Record", level=2)
    _add_table(doc,
        ["Energy Type", "Equipment/Valve ID", "Lock ID", "Locked By", "Lock Time", "Unlock Time", "Confirmed By"],
        [["", "", "", "", "", "", ""],
         ["", "", "", "", "", "", ""]],
        col_widths=[0.9, 1.0, 0.8, 0.9, 0.9, 0.9, 0.9])

    _add_heading(doc, "7.4 TBT Training Record", level=2)
    _add_table(doc,
        ["Name", "Position", "Training Content", "Date", "Signature"],
        [["", "", "", "", ""],
         ["", "", "", "", ""]],
        col_widths=[1.3, 1.0, 2.0, 1.0, 1.2])

    doc.add_page_break()

    # --- Chapter 8: Compliance Review ---
    compliance = plan_data.get("compliance_report")
    if compliance:
        _add_heading(doc, "8. Compliance Review Report", level=1)

        _add_heading(doc, "8.1 Summary", level=2)
        _add_table(doc,
            ["Item", "Result"],
            [
                ["Total Check Items", str(compliance["total_items"])],
                ["Passed", str(compliance["passed"])],
                ["Failed", str(compliance["failed"])],
                ["Pass Rate", compliance["pass_rate"]],
                ["Overall Status", compliance["overall_status"]],
            ],
            col_widths=[2.5, 4.0])

        _add_heading(doc, "8.2 Section Results", level=2)
        for section in compliance["sections"]:
            _add_heading(doc, section["regulation"], level=3)
            _add_table(doc,
                ["#", "Check Item", "Status"],
                [[str(i+1), item["item"], "PASS" if item["status"] == "pass" else "FAIL"]
                 for i, item in enumerate(section["items"])],
                col_widths=[0.4, 5.0, 1.1])

        if compliance["improvement_suggestions"]:
            _add_heading(doc, "8.3 Improvement Suggestions", level=2)
            for suggestion in compliance["improvement_suggestions"]:
                p = doc.add_paragraph()
                run = p.add_run(suggestion["section"])
                run.font.bold = True
                for missing in suggestion["missing_items"]:
                    doc.add_paragraph(missing, style="List Bullet")

    # --- Save ---
    doc.save(output_path)
    return output_path
